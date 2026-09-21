from docx import Document
from datetime import datetime

def export_to_docx(text, report, report_content, filepath, include_analysis=True):
    doc = Document()
    doc.add_heading('Writing Linter Analysis Report' if include_analysis else 'Writing Linter Document', 0)
    doc.add_paragraph(f'Date: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    
    if include_analysis:
        # 1. Stats Summary
        doc.add_heading('Executive Summary', level=1)
        summ = report.get('executive_summary', {})
        doc.add_paragraph(f"Word Count: {summ.get('word_count', 0)}")
        doc.add_paragraph(f"Est. Reading Time: {summ.get('reading_time', 'N/A')}")
        doc.add_paragraph(f"Avg Sentence Length: {float(summ.get('avg_sent_len') or 0):.1f} words")
        doc.add_paragraph(f"Dialogue Ratio: {float(summ.get('dialogue_ratio') or 0):.1f}%")
        doc.add_paragraph(f"Passive Voice: {float(summ.get('passive_perc') or 0):.1f}%")
        doc.add_paragraph(f"Prose Score: {summ.get('prose_score', 0)}/100")
        doc.add_paragraph(f"Total Issues Found: {summ.get('total_issues', 0)}")
        
        # 2. Detailed Report
        doc.add_heading('NLP Technical Report', level=1)
        doc.add_paragraph(report_content)
        
        # 3. Optimized Text Header
        doc.add_heading('Optimized Text (AI Refined)', level=1)
    
    # 3. Optimized Text
    doc.add_paragraph(text)
    
    doc.save(filepath)

def export_to_markdown(text, report, report_content, filepath, include_analysis=True):
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write("# Writing Linter Analysis Report\n\n" if include_analysis else "# Writing Linter Document\n\n")
        f.write(f"*Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n\n")
        
        if include_analysis:
            f.write("## Table of Contents\n\n")
            f.write("- [Executive Summary](#executive-summary)\n")
            f.write("- [NLP Technical Report](#nlp-technical-report)\n")
            f.write("- [AI Refined Text](#ai-refined-text)\n\n")

            f.write("## Executive Summary\n\n")
            summ = report.get('executive_summary', {})
            f.write(f"### Overall Prose Score: {summ.get('prose_score', 0)}/100\n\n")
            
            f.write("| Metric | Value |\n")
            f.write("|--------|-------|\n")
            f.write(f"| Word Count | {summ.get('word_count', 0)} |\n")
            f.write(f"| Est. Reading Time | {summ.get('reading_time', 'N/A')} |\n")
            f.write(f"| Avg Sentence Length | {float(summ.get('avg_sent_len') or 0):.1f} words |\n")
            f.write(f"| Dialogue Ratio | {float(summ.get('dialogue_ratio') or 0):.1f}% |\n")
            f.write(f"| Passive Voice | {float(summ.get('passive_perc') or 0):.1f}% |\n")
            f.write(f"| Total Issues Found | {summ.get('total_issues', 0)} |\n\n")
            
            f.write("## NLP Technical Report\n\n")
            f.write("<details>\n<summary>Click to expand technical breakdown</summary>\n\n")
            f.write("```text\n")
            f.write(report_content)
            f.write("\n```\n")
            f.write("</details>\n\n")
            
            f.write("## AI Refined Text\n\n")
        
        f.write(text)

def export_to_obsidian(text, report, report_content, filepath, include_analysis=True):
    with open(filepath, 'w', encoding='utf-8') as f:
        # YAML Frontmatter
        f.write("---\n")
        f.write("tags: [writing-linter, report]\n")
        f.write(f"date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        summ = report.get('executive_summary', {})
        f.write(f"word_count: {summ.get('word_count', 0)}\n")
        f.write("---\n\n")
 
        f.write("# Writing Linter Analysis Report\n\n" if include_analysis else "# Writing Linter Document\n\n")
        
        if include_analysis:
            f.write("> [!stats] Analysis Results\n")
            f.write(f"> - **Prose Score:** {summ.get('prose_score', 0)}/100\n")
            f.write(f"> - **Total Issues:** {summ.get('total_issues', 0)}\n")
            f.write(f"> - **Est. Reading Time:** {summ.get('reading_time', 'N/A')}\n")
            f.write(f"> - **Avg Sentence Length:** {float(summ.get('avg_sent_len') or 0):.1f} words\n")
            f.write(f"> - **Dialogue Ratio:** {float(summ.get('dialogue_ratio') or 0):.1f}%\n")
            f.write(f"> - **Passive Voice:** {float(summ.get('passive_perc') or 0):.1f}%\n\n")
            
            f.write("## NLP Technical Report\n\n")
            f.write("> [!abstract] Technical Breakdown\n")
            f.write("> <details>\n")
            f.write("> <summary>Click to expand</summary>\n")
            f.write("> \n")
            f.write("> ```text\n")
            f.write("> " + report_content.replace("\n", "\n> ") + "\n")
            f.write("> ```\n")
            f.write("> </details>\n\n")
            
            f.write("## AI Refined Text\n\n")
        
        f.write(text)

def export_to_text(text, report, report_content, filepath, include_analysis=True):
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write("Writing Linter Analysis Report\n" if include_analysis else "Writing Linter Document\n")
        f.write(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
        
        if include_analysis:
            f.write("Executive Summary\n")
            f.write("=================\n")
            summ = report.get('executive_summary', {})
            f.write(f"- Word Count: {summ.get('word_count', 0)}\n")
            f.write(f"- Est. Reading Time: {summ.get('reading_time', 'N/A')}\n")
            f.write(f"- Avg Sentence Length: {float(summ.get('avg_sent_len') or 0):.1f} words\n")
            f.write(f"- Dialogue Ratio: {float(summ.get('dialogue_ratio') or 0):.1f}%\n")
            f.write(f"- Passive Voice: {float(summ.get('passive_perc') or 0):.1f}%\n")
            f.write(f"- Prose Score: {summ.get('prose_score', 0)}/100\n")
            f.write(f"- Total Issues Found: {summ.get('total_issues', 0)}\n\n")
            
            f.write("NLP Technical Report\n")
            f.write("--------------------\n")
            f.write(report_content)
            f.write("\n\n")
            
            f.write("Optimized Text (AI Refined)\n")
            f.write("---------------------------\n")
        
        f.write(text)

import os
from jinja2 import Environment, FileSystemLoader

def export_to_html(text, report, report_content, filepath, include_analysis=True):
    # Basic HTML encoding for text
    html_text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('\n', '<br>')
    
    # Setup Jinja2 environment
    template_dir = os.path.join(os.path.dirname(__file__), 'templates')
    env = Environment(loader=FileSystemLoader(template_dir))
    
    try:
        template = env.get_template('report.html.j2')
    except Exception as e:
        # Fallback if template doesn't exist
        print(f"Template load error: {e}. Falling back to text export.")
        export_to_text(text, report, report_content, filepath, include_analysis)
        return
        
    date_str = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    summary = report.get('executive_summary', {})
    
    html_output = template.render(
        include_analysis=include_analysis,
        date_str=date_str,
        summary=summary,
        report_content=report_content,
        html_text=html_text
    )
    
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(html_output)

import json
def export_to_json(text, report, filepath, include_analysis=True):
    # Construct the enhanced JSON schema
    output = {
        "version": "1.0",
        "timestamp": datetime.now().strftime('%Y-%m-%dT%H:%M:%S'),
        "metadata": report.get("executive_summary", {}),
        "text_content": {
            "original": text,
            "fixed": report.get("fixed_text", "")
        }
    }
    
    if include_analysis:
        output["metrics"] = {
            "readability": report.get("readability", {}),
            "tone": report.get("tone", {}),
            "rhythm": report.get("paragraph_rhythm", {})
        }
        output["diagnostics"] = {
            "passive_voice": report.get("passive_voice", []),
            "pro_editing": report.get("pro_editing", []),
            "advanced_features": report.get("advanced_features", []),
            "structure": {
                "run_ons": report.get("run_ons", []),
                "clipped_sentences": report.get("clipped_sentences", []),
                "sentence_splices": report.get("sentence_splices", []),
                "dangling_participles": report.get("dangling_participles", []),
                "tense_consistency": report.get("tense_consistency", [])
            },
            "creative": {
                "nominalizations": report.get("nominalizations", []),
                "dialogue_beats": report.get("dialogue_beats", []),
                "redundancy": report.get("redundancy", []),
                "show_dont_tell": report.get("show_dont_tell", []),
                "adverbial_tags": report.get("adverbial_tags", [])
            }
        }
        if "suppressions" in report:
            output["suppressions"] = report["suppressions"]
    
    json_out = json.dumps(output, indent=2)
    if filepath:
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(json_out)
    return json_out
